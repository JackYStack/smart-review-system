from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "frontend" / "public" / "demo-assets" / "脚手架专项方案_演示审查报告.docx"
BLUE = "1677FF"
DARK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "EAF3FF"
LIGHT_GRAY = "F2F4F7"
RED = "B42318"
GOLD = "B54708"


def set_run_font(run, size=11, bold=False, color=DARK, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bottom_border(paragraph, color=BLUE, size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_issue(doc, number, severity, title, evidence, basis, gap, suggestion):
    color = RED if severity == "严重" else GOLD
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(f"{number}. [{severity}] "), size=12, bold=True, color=color)
    set_run_font(p.add_run(title), size=12, bold=True)
    for label, value in (
        ("方案证据", evidence),
        ("法规/规则依据", basis),
        ("具体差距", gap),
        ("整改建议", suggestion),
    ):
        row = doc.add_paragraph()
        row.paragraph_format.left_indent = Inches(0.18)
        row.paragraph_format.space_after = Pt(4)
        set_run_font(row.add_run(f"{label}："), bold=True, color=MUTED)
        set_run_font(row.add_run(value))


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 11, 5),
        ("Heading 3", 11.5, "344054", 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.add_run("SMARTREVIEW  |  危大工程专项方案智能审查"), size=8.5, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("演示报告  ·  第 "), size=8.5, color=MUTED)
    add_field(footer, "PAGE")
    set_run_font(footer.add_run(" 页"), size=8.5, color=MUTED)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    set_run_font(kicker.add_run("辅助审查报告  /  DEMO REVIEW REPORT"), size=9, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("落地式钢管脚手架专项施工方案"), size=23, bold=True, color="101828")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    set_run_font(subtitle.add_run("智能辅助审查结果（含缺陷演示样例）"), size=14, bold=True, color=MUTED)
    add_bottom_border(subtitle)

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2200, 7160])
    for row, (label, value) in zip(
        meta.rows,
        (
            ("任务编号", "DEMO-1001"),
            ("工程类型", "脚手架工程 / 落地式钢管脚手架"),
            ("审查状态", "发现问题，建议整改后复审"),
            ("生成时间", "2026年8月23日（演示数据）"),
        ),
    ):
        set_cell_fill(row.cells[0], LIGHT_GRAY)
        set_run_font(row.cells[0].paragraphs[0].add_run(label), bold=True, color=MUTED)
        set_run_font(row.cells[1].paragraphs[0].add_run(value))

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    set_table_geometry(callout, [9360])
    set_cell_fill(callout.cell(0, 0), LIGHT_BLUE)
    cp = callout.cell(0, 0).paragraphs[0]
    set_run_font(cp.add_run("审查结论："), size=11.5, bold=True, color=BLUE)
    set_run_font(cp.add_run("共识别5项待处理问题，其中严重问题2项、警告问题3项。建议在参数统一、地基承载力验算及现场荷载控制整改完成后，由专业人员复核。"), size=11)

    add_heading(doc, "一、审查范围与方法", 1)
    p = doc.add_paragraph()
    set_run_font(p.add_run("本报告用于展示系统输出样式，模拟文档结构检查、方案证据提取、法规知识检索、问题分级和整改建议生成。当前文件不是PaddleOCR、RAGFlow及Dify完整生产链路的真实运行结果。"))

    add_heading(doc, "二、问题汇总", 1)
    summary = doc.add_table(rows=1, cols=4)
    summary.style = "Table Grid"
    set_table_geometry(summary, [850, 1600, 4810, 2100])
    headers = ("序号", "等级", "问题", "处理建议")
    for cell, text in zip(summary.rows[0].cells, headers):
        set_cell_fill(cell, LIGHT_GRAY)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(cell.paragraphs[0].add_run(text), bold=True)
    rows = (
        ("1", "严重", "连墙件布置参数前后不一致", "统一参数并复核稳定性"),
        ("2", "严重", "立杆基础缺少承载力验算", "补充计算与排水措施"),
        ("3", "警告", "卸料平台荷载控制不可执行", "明确限载值和巡查要求"),
        ("4", "警告", "缺少施工监测专项章节", "补充监测频率和预警值"),
        ("5", "警告", "应急处置内容不完整", "补充响应流程和撤离路线"),
    )
    for item in rows:
        cells = summary.add_row().cells
        for idx, (cell, text) in enumerate(zip(cells, item)):
            if idx < 2:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(cell.paragraphs[0].add_run(text), color=RED if idx == 1 and text == "严重" else DARK)
    set_table_geometry(summary, [850, 1600, 4810, 2100])

    add_heading(doc, "三、详细审查意见", 1)
    add_issue(
        doc, "1", "严重", "连墙件布置参数前后不一致",
        "方案正文写明“连墙件按三步三跨布置”，计算与构造说明采用“两步三跨”控制。",
        "项目内部参数一致性检查；正式结论应结合现行脚手架安全技术标准原文复核。",
        "正文、计算书与构造图采用不同参数，现场交底和搭设无法形成唯一执行依据。",
        "统一正文、计算书和节点图参数；按最终布置重新复核架体稳定性，并完成审批后交底。",
    )
    add_issue(
        doc, "2", "严重", "立杆基础缺少承载力验算依据",
        "方案仅写“地面夯实后设置垫板”，未提供地基承载力取值、垫板规格和沉降控制要求。",
        "基础承载与稳定性属于脚手架方案关键控制内容；正式法规条款应由知识库召回后引用。",
        "缺少可复核的计算参数，无法判断基础是否满足架体荷载和不均匀沉降控制要求。",
        "补充地基承载力参数、荷载组合、垫板规格、排水做法和沉降巡查要求。",
    )
    add_issue(
        doc, "3", "警告", "卸料平台荷载控制措施不可执行",
        "方案规定“严禁超载”，但未给出允许荷载、限载标识和材料堆放控制方式。",
        "现场措施必须具有明确参数、责任和可检查记录。",
        "仅有原则性要求，现场人员无法据此判断何时超载，也无法形成验收依据。",
        "明确允许荷载和材料数量，悬挂限载牌，设置巡查责任人并形成每日检查记录。",
    )
    add_issue(
        doc, "4", "警告", "缺少施工监测专项章节",
        "当前方案目录未检出“施工监测”或等效章节。",
        "结构完整性规则：专项方案应明确施工期间的检查、监测和异常处置安排。",
        "缺少监测项目、频率、预警值及异常后的停工处置条件。",
        "补充架体垂直度、沉降、连接节点等监测项目，并明确频率、阈值和责任人。",
    )
    add_issue(
        doc, "5", "警告", "应急处置内容不完整",
        "应急预案仅列出联系电话，未见险情分级、响应流程、撤离路线和救援资源。",
        "结构完整性和现场可执行性规则。",
        "发生倾斜、沉降或局部失稳征兆时，无法依据方案快速组织停工、警戒和撤离。",
        "补充险情分级、响应程序、人员撤离路线、警戒范围、救援物资和外部联络机制。",
    )

    add_heading(doc, "四、人工复核与整改闭环", 1)
    for index, text in enumerate((
        "由方案编制人员逐项补充或修订，并标注修改章节。",
        "由项目技术负责人核对参数、计算书和节点图的一致性。",
        "涉及法规条款的结论应使用现行有效规范原文复核，不得仅依据本演示报告。",
        "整改完成后重新提交系统审查，并保存审查前后版本及复核记录。",
    ), start=1):
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(text))

    disclaimer = doc.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(14)
    disclaimer.paragraph_format.space_after = Pt(0)
    set_run_font(disclaimer.add_run("重要说明："), bold=True, color=RED)
    set_run_font(disclaimer.add_run("本文件为界面与导出功能演示产物，不替代法定审查、专家论证或项目技术负责人签署意见。"), color=MUTED)

    doc.core_properties.title = "脚手架专项方案演示审查报告"
    doc.core_properties.subject = "SmartReview演示导出结果"
    doc.core_properties.author = "SmartReview"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
