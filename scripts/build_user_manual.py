#!/usr/bin/env python3
"""Build docs/用户使用指导手册.docx from structured content and images."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("Missing dependency: pip install python-docx", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
IMAGES = ROOT / "images"
OUTPUT = DOCS / "用户使用指导手册.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        from docx.oxml import OxmlElement

        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_fill: str = "D9E2F3") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], header_fill)
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def add_image(doc: Document, rel_path: str, caption: str, width_cm: float = 15.5) -> None:
    img_path = IMAGES / rel_path
    if not img_path.exists():
        p = doc.add_paragraph(f"[插图缺失: {rel_path}]")
        p.runs[0].italic = True
        return
    doc.add_picture(str(img_path), width=Cm(width_cm))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run("说明：")
    run.bold = True
    p.add_run(text)


def build() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Default font for East Asian text
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # Cover
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("智能方案审核系统")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("使用指导手册")
    sr.font.size = Pt(20)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("\n版本：1.0\n适用对象：方案编制人员、系统管理员").font.size = Pt(12)

    doc.add_page_break()

    # Preface
    add_heading(doc, "前言", 1)
    doc.add_paragraph(
        "智能方案审核系统面向施工方案的全流程管理与智能审核，帮助编制人员快速完成方案上传与自动审查，"
        "也帮助管理员维护审核规则与知识库，使审核结果更准确、稳定。"
    )
    doc.add_paragraph(
        "普通用户的主要工作是：下载模版 → 本地编写 → 上传审核 → 查看审阅结果 → 导出带批注 Word。"
    )
    doc.add_paragraph(
        "管理员的主要工作是：配置大模型与知识库连接、维护方案类型与编制依据、上传 Word 模版并按 Excel 审核模版录入章节规则，"
        "以及将规范文件上传至 Dify 知识库（10.73.2.13）。"
    )
    doc.add_paragraph(
        "访问方式：使用浏览器打开系统地址（开发环境默认为 http://localhost:5173，生产环境以实际部署地址为准），"
        "使用账号或手机号与密码登录。"
    )
    add_note(
        doc,
        "系统当前不支持 Excel 一键导入审核规则，管理员需参照 Excel 模版各 Sheet，在系统中逐条手工录入。详见附录 A。",
    )
    doc.add_page_break()

    # Part 1
    add_heading(doc, "第一部分：普通用户操作指南", 1)
    doc.add_paragraph(
        "普通用户登录后，左侧菜单仅显示「方案审核」一项，所有日常操作均在该页面及关联的审阅页面完成。"
    )

    add_heading(doc, "1.1 登录系统", 2)
    add_numbered(
        doc,
        [
            "打开系统登录页。",
            "输入账号或手机号和密码。",
            "点击「立即登录」。",
            "登录成功后，普通用户将自动进入「方案审核」页面。",
        ],
    )

    add_heading(doc, "1.2 方案审核主页", 2)
    add_table(
        doc,
        ["统计项", "含义"],
        [
            ["总审核数", "当前账号可见的全部审核任务数量"],
            ["待人工审核", "状态为「已完成」的任务数（可进入人工审阅）"],
            ["审核异常", "状态为「失败」的任务数"],
        ],
    )
    add_image(doc, "方案审核.png", "图 1-1  方案审核页面")

    add_heading(doc, "1.3 下载模版", 2)
    add_numbered(
        doc,
        [
            "在「选择方案类型」下拉框中，选择与您方案对应的类型（格式为「方案大类 / 方案名称」）。下拉列表中仅显示已配置模版的方案类型。",
            "点击「下载模版」按钮。",
            "浏览器将下载 .docx 格式的 Word 模版文件。",
            "在 Microsoft Word 或 WPS 中打开模版，按章节结构填写方案内容。",
        ],
    )
    add_note(
        doc,
        "请保持 Word 文档的标题层级与模版一致（使用 Word「标题 1」至「标题 9」样式），否则上传后可能无法通过结构审核。",
    )

    add_heading(doc, "1.4 上传发起审核", 2)
    add_numbered(
        doc,
        [
            "在方案审核页，先选择方案类型（须与编写时使用的模版类型一致）。",
            "点击「方案审核」主按钮。",
            "在弹出的「提交方案审核」对话框中，确认方案类型并选择 .docx 文件。",
            "点击「提交」。系统提示「已提交审核任务」后，新任务将出现在任务列表中。",
        ],
    )

    add_heading(doc, "1.5 跟踪任务状态", 2)
    add_table(
        doc,
        ["状态", "含义"],
        [
            ["排队中", "任务已创建，等待后台处理"],
            ["处理中", "正在执行审核，可能显示当前审核阶段"],
            ["已完成", "审核流程结束，可进入人工审阅、导出文档"],
            ["失败", "审核未成功完成，仍可查看审阅报告"],
        ],
    )
    doc.add_paragraph("处理中可能显示的审核阶段：结构审核、编制依据审核、上下文一致性、内容审核、通篇审核。")

    add_heading(doc, "1.6 人工审阅", 2)
    add_numbered(
        doc,
        [
            "在任务列表中，对状态为「已完成」或「失败」的任务点击「人工审阅」。",
            "左侧为审核步骤链，右侧为审核详情与问题列表。",
            "点击左侧不同步骤，切换查看各环节的审核结果。",
        ],
    )
    add_table(
        doc,
        ["级别", "含义"],
        [
            ["严重", "需优先整改的问题"],
            ["警告", "建议关注并修改的问题"],
            ["提示", "优化类建议"],
        ],
    )
    add_image(doc, "人工审阅7.png", "图 1-2  人工审阅页面")

    add_heading(doc, "1.7 预览、编辑与导出", 2)
    add_bullets(
        doc,
        [
            "任务列表或人工审阅页点击「导出报告」/「导出 Word」，下载带批注 Word。",
            "人工审阅页点击「预览」可只读查看文档（需管理员配置 OnlyOffice）。",
            "点击「编辑」可在线修改文档；编辑后请先保存（Ctrl+S），再返回导出。",
        ],
    )
    add_image(doc, "预览.png", "图 1-3  文档预览页面")

    add_heading(doc, "1.8 常见问题", 2)
    add_table(
        doc,
        ["现象", "原因与处理"],
        [
            ["无法下载模版", "该方案类型尚未上传模版，请联系管理员"],
            ["找不到方案类型", "仅显示已配置模版的类型，请联系管理员"],
            ["任务状态为「失败」", "常见原因为结构审核未通过，请对照模版调整标题结构后重新上传"],
            ["暂无带批注文档", "任务未完成，或结构审核未通过"],
            ["预览/编辑不可用", "管理员尚未配置 OnlyOffice，或任务尚未生成输出文档"],
        ],
    )
    doc.add_page_break()

    # Part 2
    add_heading(doc, "第二部分：管理员操作指南", 1)
    doc.add_paragraph(
        "管理员登录后，左侧菜单包含：数据看板、方案类型管理、编制依据管理、模板管理、用户管理、设置等模块。"
    )

    add_heading(doc, "2.1 推荐配置顺序", 2)
    add_numbered(
        doc,
        [
            "设置 → 模型配置 — 接入大语言模型（审核必需）",
            "设置 → 知识库 — 连接 Dify（如 http://10.73.2.13/v1）并填写 API 密钥",
            "方案类型管理 — 创建「方案大类 + 方案名称」",
            "编制依据管理 — 录入编制依据库",
            "模板管理 — 上传 Word 模版 → 规则设置 → 审核工作流 →（可选）通篇审核",
            "Dify 控制台（http://10.73.2.13）— 上传知识库文档",
            "设置 → 审核配置 — 调整并发、超时、调试开关",
            "（可选）设置 → OnlyOffice — 配置在线预览与编辑",
        ],
    )

    add_heading(doc, "2.2 系统设置", 2)

    add_heading(doc, "2.2.1 知识库连接", 3)
    doc.add_paragraph("路径：设置 → 知识库")
    add_table(
        doc,
        ["配置项", "说明"],
        [
            ["Dify 服务地址", "如 http://10.73.2.13/v1"],
            ["API 密钥", "Dify 平台生成的 API Key"],
            ["知识库名称前缀过滤", "可选，用于下拉列表筛选知识库"],
        ],
    )
    add_note(
        doc,
        "SmartReview 不在本系统内上传知识库文档，仅连接 Dify 进行检索。"
        "规范原文须在 Dify 管理界面（10.73.2.13）中上传，再回到本系统绑定。",
    )
    add_image(doc, "设置dify.png", "图 2-1  Dify 知识库设置")

    add_heading(doc, "2.2.2 模型配置", 3)
    doc.add_paragraph("路径：设置 → 模型。配置 LLM API 并设置默认模型，可使用「测试连接」验证。")
    add_image(doc, "设置-模型.png", "图 2-2  模型设置")

    add_heading(doc, "2.2.3 审核配置", 3)
    doc.add_paragraph("路径：设置 → 审核配置。可调整超时、并发、提示词调试开关及系统品牌。")
    add_image(doc, "设置-系统.png", "图 2-3  审核设置")

    add_heading(doc, "2.2.4 OnlyOffice（可选）", 3)
    doc.add_paragraph("路径：设置 → OnlyOffice。配置文档服务以支持在线预览与编辑。")
    add_image(doc, "设置-office.png", "图 2-4  OnlyOffice 设置")

    add_heading(doc, "2.3 方案类型管理", 2)
    doc.add_paragraph(
        "路径：方案类型管理。新建方案类型时填写「方案大类」和「方案名称」。"
        "Excel「方案类型代码」在系统中无对应字段，请以「方案大类 + 方案名称」为准。"
    )

    add_heading(doc, "2.4 编制依据管理", 2)
    doc.add_paragraph(
        "路径：编制依据管理 → 新建编制依据。参照 Excel「编制依据库」Sheet 逐条录入。"
    )
    add_image(doc, "编制依据.png", "图 2-5  编制依据管理")

    add_heading(doc, "2.5 模板管理与规则设置", 2)
    add_image(doc, "模板管理.png", "图 2-6  模板管理")

    add_heading(doc, "2.5.1 上传 Word 模版", 3)
    add_numbered(
        doc,
        [
            "在模板管理列表中找到目标方案类型，点击「上传 Word」。",
            "选择 .docx 模版文件上传。",
            "系统自动解析 Word 标题，生成标题树。",
        ],
    )

    add_heading(doc, "2.5.2 规则设置", 3)
    add_table(
        doc,
        ["配置项", "作用"],
        [
            ["引用", "内容审核时注入其他章节文本"],
            ["知识库（Dify）", "绑定外部知识库检索规范条文"],
            ["关键字", "知识库检索关键词"],
            ["审核提示词", "章节审核指令；填写后才执行内容审核"],
            ["上下文一致性校验", "跨章节比对节点与检查说明"],
            ["编制依据审核", "开关；开启后执行编制依据校验"],
        ],
    )
    add_image(doc, "规则设置.png", "图 2-7  规则设置")

    add_heading(doc, "2.5.3 审核工作流", 3)
    doc.add_paragraph("按需开启：编制依据、上下文一致性、内容审核、通篇审核。结构审核始终执行。")
    add_image(doc, "流程开关.png", "图 2-8  审核工作流开关")

    add_heading(doc, "2.5.4 通篇审核（可选）", 3)
    doc.add_paragraph("填写通篇审核提示词，可选绑定知识库与关键字。")

    add_heading(doc, "2.6 Dify 知识库文档上传", 2)
    add_numbered(
        doc,
        [
            "浏览器访问 http://10.73.2.13，进入知识库并创建知识库（名称与 Excel 一致，如「脚手架工程」）。",
            "按 Excel「内容引用」列逐条上传规范、通知等文件。",
            "在 SmartReview 设置 → 知识库 中保存 API 地址与密钥。",
            "在模板管理 → 规则设置 中为各章节选择知识库并填写 TAG 关键字。",
        ],
    )

    add_heading(doc, "2.7 用户管理", 2)
    doc.add_paragraph("路径：用户管理。可新建用户并分配普通用户或管理员角色。")

    add_heading(doc, "2.8 数据看板", 2)
    doc.add_paragraph("路径：数据看板。汇总方案数量、任务状态、Token 消耗等运行概况。")
    add_image(doc, "数据统计.png", "图 2-9  数据看板")
    doc.add_page_break()

    # Appendix A
    add_heading(doc, "附录 A：Excel 模版 → 系统录入对照表", 1)
    doc.add_paragraph(
        "参考 Excel：专项方案审核模版_落地脚手架final.xlsx（示例：落地脚手架专项方案）。"
        "系统无 Excel 导入功能，须按下列对照表逐条手工录入。"
    )

    sections = [
        (
            "A.1 Sheet「说明与版本」",
            ["Excel 字段", "系统对应", "录入方式"],
            [
                ["模版版本", "—", "自行记录，系统不存版本号"],
                ["方案类型代码", "—", "用「方案大类 + 方案名称」代替"],
                ["方案名称", "方案类型管理", "创建方案类型"],
            ],
        ),
        (
            "A.2 Sheet「章节结构_完整性」",
            ["Excel 字段", "系统对应", "录入方式"],
            [
                ["章节编码 / 章节标题", "标题树节点", "上传 Word 模版后自动生成"],
                ["是否校验", "结构审核", "系统自动比对，无需单独开关"],
                ["完整性依据 / 备注", "—", "参考说明，不入库"],
            ],
        ),
        (
            "A.3 Sheet「编制依据库」",
            ["Excel 字段", "系统字段", "说明"],
            [
                ["文献类型", "文献类型", "下拉选择"],
                ["标准号或文号", "标准号", ""],
                ["文献名称", "文献名称", ""],
                ["版本或施行日期说明", "效力状态", "现行 / 废止 / 即将实施 等"],
                ["是否必引", "必引", "是/否"],
                ["类别", "—", "可写入备注"],
                ["分类", "方案大类 + 方案名称", "绑定方案类型"],
            ],
        ),
        (
            "A.4 Sheet「一致性规则」",
            ["Excel 字段", "系统字段", "说明"],
            [
                ["源/目标章节编码", "上下文一致性比对节点", "在标题树中按章节标题选取"],
                ["检查说明", "一致性检查说明", "文本框"],
                ["是否启用", "审核工作流总开关 + 节点配置", "未配置则跳过"],
                ["严重等级", "严重/警告/提示", "由模型输出，Excel 作参考"],
            ],
        ),
        (
            "A.5 Sheet「章节审查配置」",
            ["Excel 字段", "系统字段", "说明"],
            [
                ["章节编码", "标题树节点", "点击对应章节"],
                ["依赖章节编码", "引用", "内容审核注入引用章节"],
                ["知识库引用", "知识库 + 关键字", "选 Dify 库 + 填 tag"],
                ["人工提示词", "审核提示词", "直接粘贴"],
                ["数值审核 / 审查侧重点", "—", "可合并写入审核提示词"],
                ["编制依据", "编制依据审核开关", "需校验的章节开启"],
            ],
        ),
        (
            "A.6 Sheet「知识库」",
            ["Excel 字段", "系统对应", "录入方式"],
            [
                ["知识库名称", "Dify 数据集名称", "在 10.73.2.13 创建知识库"],
                ["TAG 信息", "规则设置 → 关键字", "如「落地脚手架」"],
                ["内容引用", "Dify 上传文档", "在 Dify 控制台上传原文"],
                ["是否启用", "Dify 状态 + 系统绑定", ""],
            ],
        ),
        (
            "A.7 Sheet「审核流程图」",
            ["Excel 内容", "系统对应"],
            [
                ["AI 上下文构造模板", "编写审核提示词的参考"],
                ["流程顺序", "审核工作流开关及环节顺序"],
            ],
        ),
    ]

    for title, headers, rows in sections:
        add_heading(doc, title, 2)
        add_table(doc, headers, rows)

    doc.add_page_break()

    add_heading(doc, "附录 B：术语表", 1)
    add_table(
        doc,
        ["术语", "含义"],
        [
            ["方案类型", "施工方案分类，由「方案大类 + 方案名称」组成"],
            ["模版", "Word 格式的标准方案结构文件"],
            ["结构审核", "比对上传文档标题树与模版是否一致"],
            ["编制依据审核", "检查必引/废止规范引用情况"],
            ["上下文一致性", "跨章节数据与表述一致性检查"],
            ["内容审核", "按章节提示词与知识库进行合规审查"],
            ["通篇审核", "全文级别综合审查"],
            ["带批注文档", "审核结果写入 Word 后的输出文件"],
        ],
    )

    add_heading(doc, "附录 C：权限说明", 1)
    add_table(
        doc,
        ["功能", "普通用户", "管理员"],
        [
            ["方案审核 / 人工审阅 / 导出", "✓", "✓"],
            ["预览 / 在线编辑", "✓", "✓"],
            ["数据看板 / 方案类型 / 模版 / 依据 / 设置 / 用户", "—", "✓"],
        ],
    )

    add_heading(doc, "附录 D：部署说明（简要）", 1)
    doc.add_paragraph(
        "系统部署、数据库迁移、Docker Compose 启动等运维事项，请参阅项目 README.md 与 backend/README.md。"
    )

    doc.save(str(OUTPUT))
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
