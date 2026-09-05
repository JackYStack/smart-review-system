"""Generate audit report Word documents from ReviewReportV1 JSON."""

from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from app.models.scheme_review_task import ReviewTaskStatus, SchemeReviewTask
from app.schemas.review_report import ReportIssue, ReportStep, ReviewReportV1
from app.services.review_pipeline import WORD_COMMENT_STEP_LABEL_CN

_STEP_ORDER = [
    "structure",
    "compilation_basis",
    "context_consistency",
    "content",
    "full_document",
    "dify_workflow",
    "rules_and_formulas",
]

_SECTION_NUM_CN = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]

_STRUCTURE_KIND_LABEL: dict[str, str] = {
    "missing_section": "缺失章节",
    "order_mismatch": "顺序不符",
    "extra_section": "多余章节",
}

_STEP_DESCRIPTION: dict[str, str] = {
    "structure": "对照模版检查必备章节是否齐全（多余章节与章节顺序不校验）。",
    "compilation_basis": "审核方案编制依据是否完整覆盖现行必引规范，并识别废止误引。",
    "context_consistency": "检查方案各章节之间参数、描述是否前后一致。",
    "content": "按模版节点逐项核查章节内容是否满足审核要求。",
    "full_document": "对全文进行通篇审核，识别跨章节或整体性风险。",
    "dify_workflow": "调用团队发布的 Dify Workflow 对原始专项方案执行独立全文审查。",
    "rules_and_formulas": "使用版本化确定性规则、参数和公式复算关键指标，并保留计算过程。",
}

_STANDARD_RE = re.compile(
    r"\b(?:GB|JGJ|DBJ|DB|CECS|T/[A-Z]+)\s*[\-/]?\s*\d{2,6}(?:\.\d+)?(?:-\d{4})?\b",
    re.IGNORECASE,
)


def _as_string_array(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(x).strip() for x in value if str(x).strip()]


def _parse_title_path_value(value: Any) -> list[str]:
    if isinstance(value, list):
        return _as_string_array(value)
    if isinstance(value, str) and value.strip():
        return [p.strip() for p in re.split(r"\s*[>＞]\s*", value) if p.strip()]
    return []


def _title_path(issue: ReportIssue) -> str:
    path = _parse_title_path_value((issue.anchor or {}).get("title_path"))
    if path:
        return " > ".join(path)
    related = issue.related or {}
    loc = related.get("location")
    if isinstance(loc, dict):
        chapter_path = _parse_title_path_value(loc.get("chapter_path"))
        if chapter_path:
            return " > ".join(chapter_path)
        chapter_text = str(loc.get("chapter_text") or "").strip()
        if chapter_text:
            return chapter_text
    if isinstance(loc, str) and loc.strip():
        parsed = _parse_title_path_value(loc)
        if parsed:
            return " > ".join(parsed)
        return loc.strip()
    chapter = _parse_title_path_value(related.get("chapter"))
    if chapter:
        return " > ".join(chapter)
    for key in ("chapter_text", "chapter_a"):
        text = str(related.get(key) or "").strip()
        if text:
            return text
    return ""


def _original_text(issue: ReportIssue) -> str:
    related = issue.related or {}
    text = str(related.get("original_text") or "").strip()
    if text:
        return text
    return str(issue.evidence or "").strip()


def _suggestions(issue: ReportIssue) -> list[str]:
    related = issue.related or {}
    out: list[str] = []
    raw = related.get("suggestions")
    if isinstance(raw, list):
        for item in raw:
            s = str(item or "").strip()
            if s:
                out.append(s)
    for key in ("suggestion", "optimize_suggestion", "optimization_suggestion"):
        s = str(related.get(key) or "").strip()
        if s:
            out.append(s)
    seen: set[str] = set()
    unique: list[str] = []
    for s in out:
        if s not in seen:
            seen.add(s)
            unique.append(s)
    return unique


def _standards(issue: ReportIssue) -> str:
    related = issue.related or {}
    doc_name = str(related.get("doc_name") or "").strip()
    standard_no = str(related.get("standard_no") or "").strip()
    if doc_name and standard_no:
        return f"《{doc_name}》{standard_no}"
    if doc_name:
        return f"《{doc_name}》"
    if standard_no:
        return standard_no
    combined = f"{issue.message} {issue.evidence}"
    hits = _STANDARD_RE.findall(combined)
    if hits:
        return "、".join(dict.fromkeys(h.replace(" ", "") for h in hits))
    return ""


def _basis_category(issue: ReportIssue) -> str:
    related = issue.related or {}
    raw = str(related.get("category") or "").strip()
    if raw in ("现行缺失", "废止误引"):
        return raw
    message = str(issue.message or "")
    if "废止" in message or "失效" in message:
        return "废止误引"
    return "现行缺失"


def _context_chapter_pair(issue: ReportIssue) -> tuple[str, str]:
    related = issue.related or {}
    chapter_from_path = _title_path(issue)
    raw_a = str(
        related.get("chapter_a") or related.get("current_chapter") or ""
    ).strip()
    path_b = _parse_title_path_value(related.get("chapter_b_path"))
    raw_b = str(
        related.get("chapter_b")
        or related.get("ref_chapter")
        or related.get("compare_chapter")
        or ""
    ).strip()
    compare = " > ".join(path_b) if path_b else (raw_b or "—")
    chapter = chapter_from_path or raw_a or "—"
    return chapter, compare


def _structure_kind_label(kind: str) -> str:
    return _STRUCTURE_KIND_LABEL.get(kind, kind or "—")


def _scheme_display_name(task: SchemeReviewTask) -> str:
    raw = (task.original_filename or "").strip() or "document"
    return re.sub(r"\.docx$", "", raw, flags=re.IGNORECASE).strip() or "document"


def _format_dt(value: datetime | None) -> str:
    if value is None:
        return "—"
    try:
        return value.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(value)


def _status_label(status: str) -> str:
    mapping = {
        ReviewTaskStatus.pending: "排队中",
        ReviewTaskStatus.processing: "处理中",
        ReviewTaskStatus.succeeded: "已完成",
        ReviewTaskStatus.failed: "失败",
        ReviewTaskStatus.canceled: "已取消",
    }
    return mapping.get(status, status)


def _conclusion_label(value: str) -> str:
    return {
        "passed": "未发现问题",
        "issues_found": "发现待处理问题",
        "not_reviewed": "无法形成结论",
    }.get(value, value or "无法形成结论")


def _completeness_label(value: str) -> str:
    return {
        "complete": "完整审查",
        "partial": "降级审查",
        "unavailable": "审查不可用",
    }.get(value, value or "审查不可用")


def _human_status_label(value: str) -> str:
    return {
        "pending": "待专家复核",
        "in_review": "专家复核中",
        "changes_requested": "待整改/复审",
        "approved": "专家已确认",
    }.get(value, value or "待专家复核")


def _evidence_status(issue: ReportIssue) -> str:
    related = issue.related or {}
    if related.get("technical_error"):
        return "技术异常，未完成判定"
    status = str(related.get("evidence_status") or "").strip()
    if status == "verified":
        return "证据链完整"
    if status == "unverified" or related.get("requires_human_judgement"):
        return "疑似问题，待人工判断"
    return "证据状态未标明，待人工核验"


def _legal_evidence(issue: ReportIssue) -> str:
    related = issue.related or {}
    standard = _standards(issue)
    clause = str(related.get("clause_no") or "").strip()
    clause_text = str(related.get("clause_text") or "").strip()
    parts = [item for item in (standard, clause, clause_text) if item]
    return "\n".join(parts) if parts else "—"


def _calculation_trace(issue: ReportIssue) -> str:
    related = issue.related or {}
    steps = related.get("steps")
    expected = related.get("expected")
    rendered: list[str] = []
    if isinstance(steps, list):
        for item in steps:
            if isinstance(item, dict):
                expression = str(item.get("expression") or "").strip()
                value = str(item.get("value") or "").strip()
                if expression or value:
                    rendered.append(f"{expression} = {value}".strip(" ="))
    if isinstance(expected, dict) and expected:
        expected_text = "，".join(
            f"{key}={value}" for key, value in expected.items() if value is not None
        )
        if expected_text:
            rendered.append(f"判定阈值：{expected_text}")
    return "\n".join(rendered)


def _set_cell_text(cell: Any, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.bold = bold


def _add_paragraph(doc: Document, text: str, *, bold: bool = False, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.bold = bold


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for col_idx, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[col_idx], header, bold=True)
    for row_idx, row in enumerate(rows):
        for col_idx, cell_text in enumerate(row):
            _set_cell_text(table.rows[row_idx + 1].cells[col_idx], cell_text)
    doc.add_paragraph()


def _add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=2)
    table.style = "Table Grid"
    _set_cell_text(table.rows[0].cells[0], "项目", bold=True)
    _set_cell_text(table.rows[0].cells[1], "内容", bold=True)
    for row_idx, (label, value) in enumerate(rows):
        _set_cell_text(table.rows[row_idx + 1].cells[0], label, bold=True)
        _set_cell_text(table.rows[row_idx + 1].cells[1], value)
    doc.add_paragraph()


def _format_suggestions(issue: ReportIssue) -> str:
    suggestions = _suggestions(issue)
    return "\n".join(suggestions) if suggestions else "—"


def _build_summary(report: ReviewReportV1, task: SchemeReviewTask) -> str:
    steps_by_id = {s.step_id: s for s in report.steps}
    parts: list[str] = []
    total_issues = 0
    for step_id in _STEP_ORDER:
        step = steps_by_id.get(step_id)
        if step is None:
            continue
        n = len(step.issues)
        if n:
            label = WORD_COMMENT_STEP_LABEL_CN.get(step_id, step_id)
            parts.append(f"{label} {n} 项")
            total_issues += n
    passed_count = sum(1 for s in report.steps if s.passed)
    step_count = len(report.steps)
    status_text = _status_label(str(task.status))
    detail = "、".join(parts) if parts else "未发现明显问题"
    conclusion = (
        "全部步骤已通过"
        if passed_count == step_count and total_issues == 0
        else f"共 {step_count} 个步骤，{passed_count} 个通过"
    )
    return (
        f"本次审核任务状态为「{status_text}」，{conclusion}。"
        f"累计发现 {total_issues} 项待关注问题（{detail}）。"
        "请结合下文明细逐项核查并整改。"
    )


def _render_structure(doc: Document, step: ReportStep) -> None:
    _add_paragraph(doc, _STEP_DESCRIPTION["structure"], style="Intense Quote")
    if not step.issues:
        _add_paragraph(doc, "本步骤无问题项。")
        return
    rows: list[list[str]] = []
    for issue in step.issues:
        kind = str((issue.related or {}).get("kind") or "")
        chapter = _title_path(issue) or "—"
        rows.append([chapter, _structure_kind_label(kind), issue.message])
    _add_table(doc, ["章节", "问题类型", "说明"], rows)


def _render_compilation_basis(doc: Document, step: ReportStep) -> None:
    _add_paragraph(doc, _STEP_DESCRIPTION["compilation_basis"], style="Intense Quote")
    if not step.issues:
        _add_paragraph(doc, "本步骤无问题项。")
        return
    rows: list[list[str]] = []
    for issue in step.issues:
        suggestions = _suggestions(issue)
        rows.append(
            [
                issue.message,
                _basis_category(issue),
                _title_path(issue) or "—",
                _original_text(issue) or "无",
                "；".join(suggestions) if suggestions else "—",
                _standards(issue) or "—",
            ]
        )
    _add_table(
        doc,
        ["问题", "类别", "章节", "原文", "修改建议", "标准依据"],
        rows,
    )


def _render_context_consistency(doc: Document, step: ReportStep) -> None:
    _add_paragraph(doc, _STEP_DESCRIPTION["context_consistency"], style="Intense Quote")
    if not step.issues:
        _add_paragraph(doc, "本步骤无问题项。")
        return
    rows: list[list[str]] = []
    for issue in step.issues:
        chapter, compare = _context_chapter_pair(issue)
        suggestions = _suggestions(issue)
        rows.append(
            [
                chapter,
                compare,
                issue.message,
                "；".join(suggestions) if suggestions else "—",
            ]
        )
    _add_table(doc, ["章节", "对比章节", "问题", "优化建议"], rows)


def _render_content_like(doc: Document, step: ReportStep) -> None:
    desc = _STEP_DESCRIPTION.get(step.step_id, "")
    if desc:
        _add_paragraph(doc, desc, style="Intense Quote")
    if not step.issues:
        _add_paragraph(doc, "本步骤无问题项。")
        return
    grouped: dict[str, list[ReportIssue]] = defaultdict(list)
    for issue in step.issues:
        key = _title_path(issue) or "未定位章节"
        grouped[key].append(issue)
    for chapter, issues in grouped.items():
        doc.add_heading(chapter, level=3)
        rows: list[list[str]] = []
        for idx, issue in enumerate(issues, start=1):
            original_or_calculation = _original_text(issue) or "无"
            calculation = _calculation_trace(issue)
            if calculation:
                original_or_calculation = f"{original_or_calculation}\n{calculation}"
            rows.append(
                [
                    str(idx),
                    _evidence_status(issue),
                    issue.message,
                    original_or_calculation,
                    _legal_evidence(issue),
                    _format_suggestions(issue),
                ]
            )
        _add_table(
            doc,
            ["序号", "证据状态", "问题", "方案原文/验算过程", "法规依据", "修改建议"],
            rows,
        )


def _render_step(doc: Document, section_num: int, step: ReportStep) -> None:
    label = WORD_COMMENT_STEP_LABEL_CN.get(step.step_id, step.step_id)
    num = _SECTION_NUM_CN[section_num - 1] if section_num <= len(_SECTION_NUM_CN) else str(section_num)
    doc.add_heading(f"{num}、{label}", level=2)
    if step.summary:
        _add_paragraph(doc, f"步骤摘要：{step.summary}")
    if step.step_id == "structure":
        _render_structure(doc, step)
    elif step.step_id == "compilation_basis":
        _render_compilation_basis(doc, step)
    elif step.step_id == "context_consistency":
        _render_context_consistency(doc, step)
    elif step.step_id in ("content", "full_document"):
        _render_content_like(doc, step)
    else:
        _render_content_like(doc, step)


def build_audit_report_docx(
    task: SchemeReviewTask,
    report: ReviewReportV1,
    *,
    system_name: str = "智能方案审核",
) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)

    scheme_name = _scheme_display_name(task)
    title = doc.add_heading(f"[{scheme_name}]审核报告", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    st = task.scheme_type
    scheme_type_text = (
        f"{st.category} / {st.name}" if st else "—"
    )
    audit_time = _format_dt(task.finished_at or task.updated_at)
    passed_count = sum(1 for s in report.steps if s.passed)
    _add_kv_table(
        doc,
        [
            ("方案类型", scheme_type_text),
            ("审核时间", audit_time),
            (
                "技术处理状态",
                f"{_status_label(str(task.status))}（{passed_count}/{len(report.steps)} 步骤通过）",
            ),
            (
                "AI审查结论",
                _conclusion_label(str(getattr(task, "review_conclusion", "not_reviewed"))),
            ),
            (
                "审查完整性",
                _completeness_label(str(getattr(task, "completeness_status", "unavailable"))),
            ),
            (
                "人工复核状态",
                _human_status_label(str(getattr(task, "human_status", "pending"))),
            ),
            ("生成系统", system_name),
            ("摘要", _build_summary(report, task)),
        ],
    )

    completeness = str(getattr(task, "completeness_status", "unavailable"))
    if completeness != "complete":
        _add_paragraph(
            doc,
            "重要提示：本次并非完整审查，存在服务失败、缺少配置或证据不足。"
            "下列结果不得作为正式合规结论，须由人工专家补充核验。",
            bold=True,
        )

    steps_by_id = {s.step_id: s for s in report.steps}
    section_num = 0
    for step_id in _STEP_ORDER:
        step = steps_by_id.get(step_id)
        if step is None:
            continue
        section_num += 1
        _render_step(doc, section_num, step)

    for step in report.steps:
        if step.step_id not in _STEP_ORDER:
            section_num += 1
            _render_step(doc, section_num, step)

    doc.add_paragraph()
    _add_paragraph(
        doc,
        "注意：本报告属于 AI 辅助审查成果。未经具备相应资格的专家逐条复核并正式签发，"
        "不得替代法定审核、专家论证或责任主体签审。",
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
