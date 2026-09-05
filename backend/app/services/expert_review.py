from __future__ import annotations

import hashlib
import io
import json
from datetime import UTC, datetime, timedelta
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from sqlalchemy.orm import Session, joinedload

from app.models.review_governance import (
    AuditEvent,
    ExpertDecision,
    ReviewIssue,
    ReviewRound,
)
from app.models.rule_engine import EvidenceSource
from app.models.scheme_review_task import ReviewTaskStatus, SchemeReviewTask
from app.models.user import User, UserRole
from app.schemas.expert_review import (
    ExpertDecisionPublic,
    ReviewIssuePublic,
    ReviewRoundDetail,
    ReviewRoundPublic,
)
from app.schemas.review_report import ReviewReportV1
from app.services.review_report_docx import build_audit_report_docx


STAFF_ROLES = {UserRole.admin, UserRole.expert, UserRole.review_admin}
REVIEW_MANAGER_ROLES = {UserRole.admin, UserRole.review_admin}
TERMINAL_TASK_STATUSES = {
    ReviewTaskStatus.succeeded,
    ReviewTaskStatus.failed,
    ReviewTaskStatus.canceled,
}
CLAIM_MINUTES = 120


def safe_json_dict(value: str | None) -> dict[str, Any]:
    try:
        parsed = json.loads(value or "{}")
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def stable_issue_key(step_id: str, raw: dict[str, Any]) -> str:
    source = str(raw.get("issue_id") or "").strip()
    if source:
        # Dify and the local review stages commonly restart issue numbering in
        # every stage (for example each stage emits ``issue-1``).  The database
        # key is unique for the whole review round, so using the upstream id by
        # itself silently dropped later issues from another stage.
        canonical_source = f"{step_id}\0{source}"
        return hashlib.sha256(canonical_source.encode("utf-8")).hexdigest()[:32]
    canonical = json.dumps(
        {
            "step_id": step_id,
            "message": str(raw.get("message") or ""),
            "evidence": str(raw.get("evidence") or ""),
            "anchor": raw.get("anchor") if isinstance(raw.get("anchor"), dict) else {},
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:32]


def _task_reviewable(task: SchemeReviewTask) -> bool:
    """Only a usable AI result may enter the expert task pool.

    A technically failed task may still contain partial JSON, but allowing it
    to be claimed and signed would turn a runtime failure into an apparently
    formal review result.
    """

    return (
        str(task.status) == str(ReviewTaskStatus.succeeded)
        and str(getattr(task, "completeness_status", "unavailable")) != "unavailable"
    )


def _utc_aware(value: datetime | None) -> datetime | None:
    """Normalize timestamps returned as naive values by MySQL drivers."""

    if value is None:
        return None
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def ensure_round_for_task(db: Session, task: SchemeReviewTask) -> ReviewRound:
    row = (
        db.query(ReviewRound)
        .options(joinedload(ReviewRound.issues))
        .filter(ReviewRound.task_id == task.id)
        .first()
    )
    if row is None:
        row = ReviewRound(
            task_id=task.id,
            round_no=1,
            status="pending" if _task_reviewable(task) else "pending_ai",
        )
        db.add(row)
        db.flush()
        record_decision(db, row, None, "round_created", None, row.status, "AI审核任务进入复核链")
        append_audit_event(db, None, "review_round", row.id, "created", {"task_id": task.id})
    elif row.status == "pending_ai" and _task_reviewable(task):
        previous = row.status
        row.status = "pending"
        record_decision(db, row, None, "ai_review_finished", previous, row.status, "")

    materialize_issues(db, row, task.review_result_json)
    return row


def materialize_issues(db: Session, review_round: ReviewRound, raw_report: str | None) -> int:
    report = safe_json_dict(raw_report)
    steps = report.get("steps")
    if not isinstance(steps, list):
        return 0
    existing = {it.issue_key for it in review_round.issues}
    added = 0
    for step in steps:
        if not isinstance(step, dict):
            continue
        step_id = str(step.get("step_id") or "unknown")[:64]
        issues = step.get("issues")
        if not isinstance(issues, list):
            continue
        for raw in issues:
            if not isinstance(raw, dict):
                continue
            key = stable_issue_key(step_id, raw)
            if key in existing:
                continue
            anchor = raw.get("anchor") if isinstance(raw.get("anchor"), dict) else {}
            related = raw.get("related") if isinstance(raw.get("related"), dict) else {}
            issue_row = ReviewIssue(
                    review_round_id=review_round.id,
                    issue_key=key,
                    source_issue_id=str(raw.get("issue_id") or "")[:64],
                    step_id=step_id,
                    severity=str(raw.get("severity") or "warning")[:16],
                    message=str(raw.get("message") or "待复核问题"),
                    evidence=str(raw.get("evidence") or ""),
                    anchor_json=json.dumps(anchor, ensure_ascii=False, separators=(",", ":")),
                    related_json=json.dumps(related, ensure_ascii=False, separators=(",", ":")),
                )
            db.add(issue_row)
            db.flush()
            standard_no = str(related.get("standard_no") or "").strip()
            clause_no = str(related.get("clause_no") or "").strip()
            clause_text = str(related.get("clause_text") or "").strip()
            if standard_no and clause_no and clause_text:
                score_raw = related.get("retrieval_score")
                try:
                    retrieval_score = float(score_raw) if score_raw is not None else None
                except (TypeError, ValueError):
                    retrieval_score = None
                db.add(
                    EvidenceSource(
                        task_id=review_round.task_id,
                        review_issue_id=issue_row.id,
                        source_kind="regulation",
                        standard_no=standard_no,
                        standard_name=str(related.get("standard_name") or ""),
                        standard_version=str(related.get("standard_version") or ""),
                        effect_status=str(related.get("effect_status") or ""),
                        clause_no=clause_no,
                        clause_text=clause_text,
                        page_no=str(related.get("page_no") or ""),
                        dataset_id=str(related.get("dataset_id") or ""),
                        segment_id=str(related.get("segment_id") or ""),
                        retrieval_score=retrieval_score,
                        scheme_quote=str(raw.get("evidence") or ""),
                        location_json=json.dumps(
                            anchor, ensure_ascii=False, separators=(",", ":")
                        ),
                    )
                )
            existing.add(key)
            added += 1
    if added:
        db.flush()
    return added


def append_audit_event(
    db: Session,
    actor_id: int | None,
    entity_type: str,
    entity_id: int | str,
    action: str,
    data: dict[str, Any] | None = None,
) -> None:
    db.add(
        AuditEvent(
            actor_id=actor_id,
            entity_type=entity_type[:64],
            entity_id=str(entity_id)[:64],
            action=action[:64],
            data_json=json.dumps(data or {}, ensure_ascii=False, separators=(",", ":")),
        )
    )


def record_decision(
    db: Session,
    review_round: ReviewRound,
    actor_id: int | None,
    action: str,
    from_status: str | None,
    to_status: str | None,
    comment: str,
    *,
    issue_id: int | None = None,
) -> None:
    db.add(
        ExpertDecision(
            review_round_id=review_round.id,
            issue_id=issue_id,
            actor_id=actor_id,
            action=action,
            from_status=from_status,
            to_status=to_status,
            comment=comment,
        )
    )


def assert_staff(user: User) -> None:
    if user.role not in STAFF_ROLES:
        raise PermissionError("需要专家或复核管理权限")


def assert_round_editor(user: User, review_round: ReviewRound) -> None:
    assert_staff(user)
    if user.role in REVIEW_MANAGER_ROLES:
        return
    if review_round.assigned_expert_id != user.id:
        raise PermissionError("请先领取该复核任务")


def claim_round(db: Session, review_round: ReviewRound, user: User) -> ReviewRound:
    assert_staff(user)
    now = datetime.now(UTC)
    if not _task_reviewable(review_round.task):
        raise ValueError("AI审核尚未形成可复核结果")
    if review_round.status not in ("pending", "in_review"):
        raise ValueError("当前状态不可领取")
    claim_expires_at = _utc_aware(review_round.claim_expires_at)
    if (
        review_round.assigned_expert_id
        and review_round.assigned_expert_id != user.id
        and (claim_expires_at is None or claim_expires_at > now)
    ):
        raise ValueError("任务已被其他专家领取")
    previous = review_round.status
    review_round.status = "in_review"
    review_round.assigned_expert_id = user.id
    review_round.claimed_at = now
    review_round.claim_expires_at = now + timedelta(minutes=CLAIM_MINUTES)
    if hasattr(review_round.task, "human_status"):
        review_round.task.human_status = "in_review"
    record_decision(db, review_round, user.id, "claimed", previous, review_round.status, "")
    append_audit_event(db, user.id, "review_round", review_round.id, "claimed", {})
    return review_round


def release_round(db: Session, review_round: ReviewRound, user: User, comment: str = "") -> ReviewRound:
    assert_round_editor(user, review_round)
    if review_round.status != "in_review" or review_round.assigned_expert_id is None:
        raise ValueError("仅已领取且审阅中的任务可释放")
    previous = review_round.status
    review_round.status = "pending"
    review_round.assigned_expert_id = None
    review_round.claimed_at = None
    review_round.claim_expires_at = None
    if hasattr(review_round.task, "human_status"):
        review_round.task.human_status = "pending"
    record_decision(db, review_round, user.id, "released", previous, review_round.status, comment)
    append_audit_event(db, user.id, "review_round", review_round.id, "released", {"comment": comment})
    return review_round


def update_issue_decision(
    db: Session,
    review_round: ReviewRound,
    issue: ReviewIssue,
    user: User,
    *,
    disposition: str,
    reviewer_comment: str,
    final_severity: str | None,
) -> ReviewIssue:
    assert_round_editor(user, review_round)
    if review_round.status != "in_review":
        raise ValueError("复核任务不在审阅中")
    previous = issue.disposition
    issue.disposition = disposition
    issue.reviewer_comment = reviewer_comment
    issue.final_severity = final_severity
    issue.reviewed_by_id = user.id
    issue.reviewed_at = datetime.now(UTC)
    record_decision(
        db,
        review_round,
        user.id,
        "issue_decided",
        previous,
        disposition,
        reviewer_comment,
        issue_id=issue.id,
    )
    append_audit_event(
        db,
        user.id,
        "review_issue",
        issue.id,
        "decided",
        {"from": previous, "to": disposition, "final_severity": final_severity},
    )
    return issue


def unresolved_issue_count(review_round: ReviewRound) -> int:
    return sum(1 for it in review_round.issues if it.disposition == "pending")


def transition_round(
    db: Session,
    review_round: ReviewRound,
    user: User,
    *,
    action: str,
    comment: str,
    conclusion: str | None = None,
) -> ReviewRound:
    assert_round_editor(user, review_round)
    previous = review_round.status
    now = datetime.now(UTC)
    if action == "request_changes":
        if previous != "in_review":
            raise ValueError("仅审阅中的任务可退回整改")
        if not comment.strip():
            raise ValueError("退回整改必须填写意见")
        review_round.status = "changes_requested"
        review_round.conclusion = conclusion or "changes_required"
        if hasattr(review_round.task, "human_status"):
            review_round.task.human_status = "changes_requested"
    elif action == "approve":
        if previous != "in_review":
            raise ValueError("仅审阅中的任务可批准")
        if not _task_reviewable(review_round.task):
            raise ValueError("AI审核结果不可用，不能批准")
        if unresolved_issue_count(review_round):
            raise ValueError("仍有未处置问题，不能批准")
        review_round.status = "approved"
        review_round.conclusion = conclusion or "passed"
        review_round.approved_at = now
        if hasattr(review_round.task, "human_status"):
            review_round.task.human_status = "approved"
    else:
        raise ValueError("不支持的复核动作")
    review_round.final_comment = comment
    review_round.claim_expires_at = None
    record_decision(db, review_round, user.id, action, previous, review_round.status, comment)
    append_audit_event(
        db,
        user.id,
        "review_round",
        review_round.id,
        action,
        {"from": previous, "to": review_round.status, "conclusion": review_round.conclusion},
    )
    return review_round


def _issue_public(issue: ReviewIssue) -> ReviewIssuePublic:
    return ReviewIssuePublic(
        id=issue.id,
        review_round_id=issue.review_round_id,
        issue_key=issue.issue_key,
        source_issue_id=issue.source_issue_id,
        step_id=issue.step_id,
        severity=issue.severity,
        message=issue.message,
        evidence=issue.evidence,
        anchor=safe_json_dict(issue.anchor_json),
        related=safe_json_dict(issue.related_json),
        disposition=issue.disposition,
        reviewer_comment=issue.reviewer_comment,
        final_severity=issue.final_severity,
        reviewed_by_id=issue.reviewed_by_id,
        reviewed_at=issue.reviewed_at,
        created_at=issue.created_at,
        updated_at=issue.updated_at,
    )


def round_public(review_round: ReviewRound, *, detail: bool = False) -> ReviewRoundPublic:
    task = review_round.task
    scheme = task.scheme_type if task else None
    owner = task.user if task else None
    revision = review_round.document_revision
    project = revision.project if revision else None
    base: dict[str, Any] = {
        "id": review_round.id,
        "task_id": review_round.task_id,
        "document_revision_id": review_round.document_revision_id,
        "project_id": project.id if project else None,
        "project_name": project.name if project else "",
        "project_region": project.region if project else "",
        "version_label": revision.version_label if revision else "",
        "parent_round_id": review_round.parent_round_id,
        "round_no": review_round.round_no,
        "status": review_round.status,
        "assigned_expert_id": review_round.assigned_expert_id,
        "assigned_expert_username": (
            review_round.assigned_expert.username if review_round.assigned_expert else None
        ),
        "claimed_at": review_round.claimed_at,
        "claim_expires_at": review_round.claim_expires_at,
        "conclusion": review_round.conclusion,
        "final_comment": review_round.final_comment,
        "approved_at": review_round.approved_at,
        "signed_by_id": review_round.signed_by_id,
        "signed_at": review_round.signed_at,
        "signed_report_available": bool(review_round.signed_report_object_key),
        "scheme_category": scheme.category if scheme else "",
        "scheme_name": scheme.name if scheme else "",
        "original_filename": task.original_filename if task else "",
        "task_owner_username": owner.username if owner else "",
        "review_conclusion": str(getattr(task, "review_conclusion", "not_reviewed")),
        "completeness_status": str(getattr(task, "completeness_status", "unavailable")),
        "issue_count": len(review_round.issues),
        "pending_issue_count": unresolved_issue_count(review_round),
        "created_at": review_round.created_at,
        "updated_at": review_round.updated_at,
    }
    if not detail:
        return ReviewRoundPublic(**base)
    base["issues"] = [_issue_public(it) for it in review_round.issues]
    base["decisions"] = [ExpertDecisionPublic.model_validate(it) for it in review_round.decisions]
    return ReviewRoundDetail(**base)


def build_signed_report(review_round: ReviewRound, *, system_name: str) -> tuple[bytes, str]:
    task = review_round.task
    raw = safe_json_dict(task.review_result_json)
    report = ReviewReportV1.model_validate(raw)
    ai_bytes = build_audit_report_docx(task, report, system_name=system_name)
    doc = Document(io.BytesIO(ai_bytes))
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(10.5)
    doc.add_page_break()
    doc.add_heading("专家复核与签发结论", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    reviewer = review_round.assigned_expert.username if review_round.assigned_expert else "—"
    signer = review_round.signed_by.username if review_round.signed_by else "—"
    for label, value in (
        ("复核结论", review_round.conclusion or "—"),
        ("复核专家", reviewer),
        ("签发人", signer),
        ("复核意见", review_round.final_comment or "—"),
        ("问题处置", f"共 {len(review_round.issues)} 项，未处置 {unresolved_issue_count(review_round)} 项"),
    ):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    doc.add_heading("逐项复核记录", level=2)
    issue_table = doc.add_table(rows=1, cols=5)
    issue_table.style = "Table Grid"
    for idx, value in enumerate(("序号", "AI问题", "专家处置", "最终等级", "专家意见")):
        issue_table.rows[0].cells[idx].text = value
    for idx, issue in enumerate(review_round.issues, start=1):
        cells = issue_table.add_row().cells
        cells[0].text = str(idx)
        cells[1].text = issue.message
        cells[2].text = issue.disposition
        cells[3].text = issue.final_severity or issue.severity
        cells[4].text = issue.reviewer_comment or "—"
    out = io.BytesIO()
    doc.save(out)
    content = out.getvalue()
    return content, hashlib.sha256(content).hexdigest()
