"""Validated Excel ETL for deterministic review rules and formulas.

The import is intentionally two-phase: parse/validate first, then apply the
entire workbook in one database transaction.  A bad row never produces a
partially imported rule set.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.models.rule_engine import FormulaDefinition, RuleDefinition
from app.models.scheme_type import SchemeType
from app.schemas.rule_engine import FormulaDefinitionCreate, RuleDefinitionCreate


RULE_SHEET_NAMES = ("rules", "规则", "规则定义")
FORMULA_SHEET_NAMES = ("formulas", "公式", "公式定义")

HEADER_ALIASES = {
    "规则编号": "rule_code",
    "规则名称": "name",
    "规则类型": "rule_type",
    "版本": "version",
    "严重等级": "severity",
    "是否启用": "enabled",
    "规则配置": "config_json",
    "规范编号": "source_standard_no",
    "规范名称": "source_standard_name",
    "规范版本": "source_version",
    "条款号": "source_clause",
    "条款原文": "source_text",
    "公式编号": "formula_code",
    "公式名称": "formula_name",
    "规则版本": "rule_version",
    "表达式": "expression",
    "变量映射": "variables_json",
    "结果单位": "result_unit",
    "比较符": "comparator",
    "阈值": "threshold_value",
    "阈值参数": "threshold_parameter",
}


@dataclass(frozen=True)
class RowError:
    sheet: str
    row: int
    field: str
    message: str


@dataclass
class RuleImportPlan:
    scheme_type_id: int
    rules: list[RuleDefinitionCreate] = field(default_factory=list)
    formulas: list[dict[str, Any]] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.errors

    def public_report(self) -> dict[str, Any]:
        return {
            "valid": self.valid,
            "scheme_type_id": self.scheme_type_id,
            "rule_count": len(self.rules),
            "formula_count": len(self.formulas),
            "errors": [error.__dict__ for error in self.errors],
        }


def _sheet(workbook, names: tuple[str, ...]):
    lookup = {name.strip().lower(): name for name in workbook.sheetnames}
    for candidate in names:
        real = lookup.get(candidate.lower())
        if real:
            return workbook[real]
    return None


def _normalize_header(value: object) -> str:
    text = str(value or "").strip()
    return HEADER_ALIASES.get(text, text.lower())


def _rows(sheet) -> list[tuple[int, dict[str, Any]]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return []
    headers = [_normalize_header(value) for value in raw_headers]
    result: list[tuple[int, dict[str, Any]]] = []
    for row_no, values in enumerate(iterator, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        result.append(
            (
                row_no,
                {
                    header: value
                    for header, value in zip(headers, values, strict=False)
                    if header
                },
            )
        )
    return result


def _bool(value: object, default: bool = True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "是", "启用"}:
        return True
    if normalized in {"0", "false", "no", "n", "否", "停用"}:
        return False
    raise ValueError(f"无法识别布尔值: {value}")


def _json_object(value: object, *, default: dict[str, Any] | None = None) -> dict[str, Any]:
    if value is None or str(value).strip() == "":
        return dict(default or {})
    if isinstance(value, dict):
        return value
    parsed = json.loads(str(value))
    if not isinstance(parsed, dict):
        raise ValueError("必须是JSON对象")
    return parsed


def parse_rule_workbook(path: str | Path, *, scheme_type_id: int) -> RuleImportPlan:
    source = Path(path)
    if not source.is_file() or source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("规则文件必须是存在的 .xlsx 或 .xlsm 文件")
    workbook = load_workbook(source, read_only=True, data_only=True)
    plan = RuleImportPlan(scheme_type_id=scheme_type_id)
    rules_sheet = _sheet(workbook, RULE_SHEET_NAMES)
    formulas_sheet = _sheet(workbook, FORMULA_SHEET_NAMES)
    if rules_sheet is None:
        plan.errors.append(RowError("rules", 1, "sheet", "缺少 rules（规则）工作表"))
        return plan

    seen_rules: set[tuple[str, int]] = set()
    for row_no, raw in _rows(rules_sheet):
        try:
            payload = RuleDefinitionCreate(
                scheme_type_id=scheme_type_id,
                rule_code=str(raw.get("rule_code") or "").strip(),
                version=int(raw.get("version") or 1),
                name=str(raw.get("name") or "").strip(),
                rule_type=str(raw.get("rule_type") or "").strip(),
                severity=str(raw.get("severity") or "error").strip().lower(),
                enabled=_bool(raw.get("enabled")),
                config=_json_object(raw.get("config_json")),
                source_standard_no=str(raw.get("source_standard_no") or "").strip(),
                source_standard_name=str(raw.get("source_standard_name") or "").strip(),
                source_version=str(raw.get("source_version") or "").strip(),
                source_clause=str(raw.get("source_clause") or "").strip(),
                source_text=str(raw.get("source_text") or "").strip(),
            )
            key = (payload.rule_code, payload.version)
            if key in seen_rules:
                raise ValueError(f"规则编号与版本重复: {key[0]} v{key[1]}")
            if not (
                payload.source_standard_no
                and payload.source_clause
                and payload.source_text
            ):
                raise ValueError("缺少规范编号、条款号或条款原文，不允许形成确定性规则")
            seen_rules.add(key)
            plan.rules.append(payload)
        except (ValueError, TypeError, json.JSONDecodeError, ValidationError) as exc:
            plan.errors.append(RowError(rules_sheet.title, row_no, "row", str(exc)))

    if formulas_sheet is not None:
        seen_formulas: set[tuple[str, str, int]] = set()
        for row_no, raw in _rows(formulas_sheet):
            try:
                rule_code = str(raw.get("rule_code") or "").strip()
                rule_version = int(raw.get("rule_version") or 1)
                formula_version = int(raw.get("version") or 1)
                formula_code = str(raw.get("formula_code") or "").strip()
                key = (rule_code, formula_code, formula_version)
                if key in seen_formulas:
                    raise ValueError("同一规则下公式编号与版本重复")
                if (rule_code, rule_version) not in seen_rules:
                    raise ValueError(f"找不到本工作簿规则: {rule_code} v{rule_version}")
                threshold_raw = raw.get("threshold_value")
                parsed = FormulaDefinitionCreate(
                    rule_id=1,  # apply阶段解析真实外键
                    formula_code=formula_code,
                    version=formula_version,
                    name=str(raw.get("formula_name") or raw.get("name") or "").strip(),
                    expression=str(raw.get("expression") or "").strip(),
                    variables=_json_object(raw.get("variables_json")),
                    result_unit=str(raw.get("result_unit") or "").strip(),
                    comparator=str(raw.get("comparator") or "<=").strip(),
                    threshold_value=(
                        Decimal(str(threshold_raw))
                        if threshold_raw is not None and str(threshold_raw).strip()
                        else None
                    ),
                    threshold_parameter=(
                        str(raw.get("threshold_parameter") or "").strip() or None
                    ),
                    enabled=_bool(raw.get("enabled")),
                )
                if parsed.threshold_value is None and not parsed.threshold_parameter:
                    raise ValueError("阈值或阈值参数至少填写一个")
                seen_formulas.add(key)
                plan.formulas.append(
                    {
                        "rule_code": rule_code,
                        "rule_version": rule_version,
                        "definition": parsed,
                    }
                )
            except (ValueError, TypeError, json.JSONDecodeError, ValidationError) as exc:
                plan.errors.append(RowError(formulas_sheet.title, row_no, "row", str(exc)))
    return plan


def apply_rule_import(
    db: Session,
    plan: RuleImportPlan,
    *,
    actor_id: int | None,
    replace_existing: bool = False,
) -> dict[str, int]:
    if not plan.valid:
        raise ValueError("导入计划存在校验错误，禁止写入数据库")
    if db.get(SchemeType, plan.scheme_type_id) is None:
        raise ValueError("方案类型不存在")
    rule_map: dict[tuple[str, int], RuleDefinition] = {}
    created_rules = 0
    created_formulas = 0
    for item in plan.rules:
        row = (
            db.query(RuleDefinition)
            .filter(
                RuleDefinition.scheme_type_id == plan.scheme_type_id,
                RuleDefinition.rule_code == item.rule_code,
                RuleDefinition.version == item.version,
            )
            .first()
        )
        if row is not None and not replace_existing:
            raise ValueError(f"规则已存在: {item.rule_code} v{item.version}")
        if row is None:
            row = RuleDefinition(
                scheme_type_id=plan.scheme_type_id,
                rule_code=item.rule_code,
                version=item.version,
                created_by_id=actor_id,
            )
            db.add(row)
            created_rules += 1
        row.name = item.name
        row.rule_type = item.rule_type
        row.severity = item.severity
        row.enabled = item.enabled
        row.config_json = json.dumps(item.config, ensure_ascii=False, separators=(",", ":"))
        row.source_standard_no = item.source_standard_no
        row.source_standard_name = item.source_standard_name
        row.source_version = item.source_version
        row.source_clause = item.source_clause
        row.source_text = item.source_text
        db.flush()
        rule_map[(item.rule_code, item.version)] = row

    for raw in plan.formulas:
        definition: FormulaDefinitionCreate = raw["definition"]
        rule = rule_map[(raw["rule_code"], raw["rule_version"])]
        row = (
            db.query(FormulaDefinition)
            .filter(
                FormulaDefinition.rule_id == rule.id,
                FormulaDefinition.formula_code == definition.formula_code,
                FormulaDefinition.version == definition.version,
            )
            .first()
        )
        if row is not None and not replace_existing:
            raise ValueError(
                f"公式已存在: {definition.formula_code} v{definition.version}"
            )
        if row is None:
            row = FormulaDefinition(
                rule_id=rule.id,
                formula_code=definition.formula_code,
                version=definition.version,
            )
            db.add(row)
            created_formulas += 1
        row.name = definition.name
        row.expression = definition.expression
        row.variables_json = json.dumps(
            definition.variables, ensure_ascii=False, separators=(",", ":")
        )
        row.result_unit = definition.result_unit
        row.comparator = definition.comparator
        row.threshold_value = definition.threshold_value
        row.threshold_parameter = definition.threshold_parameter
        row.enabled = definition.enabled
    db.flush()
    scheme = db.get(SchemeType, plan.scheme_type_id)
    if scheme is not None:
        scheme.lifecycle_status = (
            "pending_validation" if scheme.lifecycle_status == "published" else "draft"
        )
        scheme.published_at = None
        scheme.published_by_id = None
        scheme.published_template_version_id = None
    return {"rules_created": created_rules, "formulas_created": created_formulas}
