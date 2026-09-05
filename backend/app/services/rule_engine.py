from __future__ import annotations

import ast
import io
import json
import operator
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Callable

from docx import Document
from sqlalchemy.orm import Session

from app.models.rule_engine import (
    CalculationResult,
    FormulaDefinition,
    ParameterValue,
    RuleDefinition,
)
from app.models.scheme_review_task import SchemeReviewTask
from app.services import minio_storage


class RuleExecutionError(ValueError):
    pass


@dataclass(frozen=True)
class UnitSpec:
    dimension: str
    canonical: str
    multiplier: Decimal


UNIT_SPECS: dict[str, UnitSpec] = {
    "": UnitSpec("scalar", "", Decimal("1")),
    "%": UnitSpec("ratio", "ratio", Decimal("0.01")),
    "ratio": UnitSpec("ratio", "ratio", Decimal("1")),
    "mm": UnitSpec("length", "m", Decimal("0.001")),
    "cm": UnitSpec("length", "m", Decimal("0.01")),
    "m": UnitSpec("length", "m", Decimal("1")),
    "mm2": UnitSpec("area", "m2", Decimal("0.000001")),
    "cm2": UnitSpec("area", "m2", Decimal("0.0001")),
    "m2": UnitSpec("area", "m2", Decimal("1")),
    "n": UnitSpec("force", "N", Decimal("1")),
    "kn": UnitSpec("force", "N", Decimal("1000")),
    "pa": UnitSpec("pressure", "Pa", Decimal("1")),
    "kpa": UnitSpec("pressure", "Pa", Decimal("1000")),
    "mpa": UnitSpec("pressure", "Pa", Decimal("1000000")),
    "kg": UnitSpec("mass", "kg", Decimal("1")),
    "t": UnitSpec("mass", "kg", Decimal("1000")),
}


def _unit_key(unit: str | None) -> str:
    return (unit or "").strip().lower().replace("²", "2").replace("^2", "2")


def normalize_value(value: Decimal | int | float | str, unit: str | None) -> tuple[Decimal, str]:
    try:
        number = value if isinstance(value, Decimal) else Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise RuleExecutionError(f"无法解析数值: {value}") from exc
    key = _unit_key(unit)
    spec = UNIT_SPECS.get(key)
    if spec is None:
        raise RuleExecutionError(f"不支持的单位: {unit}")
    return number * spec.multiplier, spec.canonical


def convert_value(
    value: Decimal | int | float | str, from_unit: str | None, to_unit: str | None
) -> Decimal:
    normalized, canonical = normalize_value(value, from_unit)
    target = UNIT_SPECS.get(_unit_key(to_unit))
    if target is None:
        raise RuleExecutionError(f"不支持的目标单位: {to_unit}")
    if target.canonical != canonical:
        raise RuleExecutionError(f"单位维度不一致: {from_unit} -> {to_unit}")
    return normalized / target.multiplier


def _decimal_constant(value: object) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise RuleExecutionError("公式只允许数值常量")
    return Decimal(str(value))


_BIN_OPS: dict[type[ast.operator], Callable[[Decimal, Decimal], Decimal]] = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.Mod: operator.mod,
}


def safe_evaluate(expression: str, variables: dict[str, Decimal]) -> tuple[Decimal, list[dict[str, str]]]:
    if len(expression) > 4000:
        raise RuleExecutionError("公式过长")
    try:
        tree = ast.parse(expression, mode="eval")
    except SyntaxError as exc:
        raise RuleExecutionError("公式语法错误") from exc
    steps: list[dict[str, str]] = []

    def visit(node: ast.AST) -> Decimal:
        if isinstance(node, ast.Expression):
            return visit(node.body)
        if isinstance(node, ast.Constant):
            return _decimal_constant(node.value)
        if isinstance(node, ast.Name):
            if node.id not in variables:
                raise RuleExecutionError(f"缺少公式变量: {node.id}")
            return variables[node.id]
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = visit(node.operand)
            return value if isinstance(node.op, ast.UAdd) else -value
        if isinstance(node, ast.BinOp):
            left = visit(node.left)
            right = visit(node.right)
            if isinstance(node.op, ast.Pow):
                if right != right.to_integral_value() or abs(right) > 20:
                    raise RuleExecutionError("幂指数必须是绝对值不超过20的整数")
                result = left ** int(right)
            else:
                fn = _BIN_OPS.get(type(node.op))
                if fn is None:
                    raise RuleExecutionError("公式包含不允许的运算符")
                try:
                    result = fn(left, right)
                except (ArithmeticError, InvalidOperation, ZeroDivisionError) as exc:
                    raise RuleExecutionError(f"公式运算失败: {exc!s}") from exc
            steps.append({"expression": ast.unparse(node), "value": str(result)})
            return result
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            name = node.func.id
            args = [visit(arg) for arg in node.args]
            if node.keywords:
                raise RuleExecutionError("公式函数不允许关键字参数")
            if name == "abs" and len(args) == 1:
                return abs(args[0])
            if name == "min" and args:
                return min(args)
            if name == "max" and args:
                return max(args)
            if name == "sqrt" and len(args) == 1:
                if args[0] < 0:
                    raise RuleExecutionError("负数不能开平方")
                return args[0].sqrt()
            raise RuleExecutionError(f"公式函数不允许: {name}")
        raise RuleExecutionError(f"公式包含不安全语法: {type(node).__name__}")

    result = visit(tree)
    if len(list(ast.walk(tree))) > 200:
        raise RuleExecutionError("公式复杂度超过限制")
    return result, steps


def compare_values(left: Decimal, comparator: str, right: Decimal) -> bool:
    mapping: dict[str, Callable[[Decimal, Decimal], bool]] = {
        "<": operator.lt,
        "<=": operator.le,
        "==": operator.eq,
        ">=": operator.ge,
        ">": operator.gt,
        "!=": operator.ne,
    }
    fn = mapping.get(comparator)
    if fn is None:
        raise RuleExecutionError(f"不支持的比较符: {comparator}")
    return fn(left, right)


def _json_dict(raw: str | None) -> dict[str, Any]:
    try:
        value = json.loads(raw or "{}")
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return value if isinstance(value, dict) else {}


def extract_docx_text(task: SchemeReviewTask) -> str:
    data = minio_storage.get_object_bytes(task.object_key)
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(paragraphs)


def parameter_map(rows: list[ParameterValue]) -> dict[str, ParameterValue]:
    return {row.parameter_name: row for row in rows}


def _parameter_evidence(
    row: ParameterValue,
    *,
    value: Decimal,
    unit: str,
) -> dict[str, Any]:
    """Persist the value and its provenance used by a deterministic result."""

    return {
        "parameter": row.parameter_name,
        "value": str(value),
        "unit": unit,
        "raw_value": row.raw_value,
        "source": _json_dict(row.source_json),
        "extraction_method": row.extraction_method,
        "confidence": str(row.confidence) if row.confidence is not None else None,
        "verified": row.verified_by_id is not None and row.verified_at is not None,
        "verified_by_id": row.verified_by_id,
    }


def _parameter_number(row: ParameterValue, target_unit: str | None = None) -> Decimal:
    if row.numeric_value is None:
        raise RuleExecutionError(f"参数 {row.parameter_name} 没有数值")
    if target_unit is None:
        if row.normalized_value is not None:
            return Decimal(row.normalized_value)
        return Decimal(row.numeric_value)
    return convert_value(Decimal(row.numeric_value), row.unit, target_unit)


def _evaluate_required_section(rule: RuleDefinition, config: dict[str, Any], text: str) -> tuple[bool, str, dict]:
    keywords = [str(x).strip() for x in config.get("keywords", []) if str(x).strip()]
    if not keywords:
        raise RuleExecutionError("必备章节规则未配置keywords")
    hits = [keyword for keyword in keywords if keyword in text]
    mode = str(config.get("match") or "all")
    passed = bool(hits) if mode == "any" else len(hits) == len(keywords)
    missing = [x for x in keywords if x not in hits]
    message = "已检出必备章节" if passed else f"缺少必备章节/关键词: {'、'.join(missing)}"
    return passed, message, {"keywords": keywords, "hits": hits, "match": mode}


def _evaluate_parameter_range(
    rule: RuleDefinition, config: dict[str, Any], params: dict[str, ParameterValue]
) -> tuple[bool, Decimal, str, dict, dict]:
    name = str(config.get("parameter") or "").strip()
    if not name or name not in params:
        raise RuleExecutionError(f"缺少参数: {name or '-'}")
    unit = str(config.get("unit") or params[name].unit)
    value = _parameter_number(params[name], unit)
    minimum = Decimal(str(config["min"])) if config.get("min") is not None else None
    maximum = Decimal(str(config["max"])) if config.get("max") is not None else None
    inclusive = bool(config.get("inclusive", True))
    passed = True
    if minimum is not None:
        passed = passed and (value >= minimum if inclusive else value > minimum)
    if maximum is not None:
        passed = passed and (value <= maximum if inclusive else value < maximum)
    expected = {"min": str(minimum) if minimum is not None else None, "max": str(maximum) if maximum is not None else None, "unit": unit, "inclusive": inclusive}
    inputs = {name: _parameter_evidence(params[name], value=value, unit=unit)}
    message = f"{name}={value}{unit}，{'符合' if passed else '不符合'}允许范围"
    return passed, value, unit, inputs, expected | {"message": message}


def _evaluate_cross_field(
    config: dict[str, Any], params: dict[str, ParameterValue]
) -> tuple[bool, str, dict, dict]:
    left_name = str(config.get("left") or "").strip()
    right_name = str(config.get("right") or "").strip()
    if left_name not in params or right_name not in params:
        raise RuleExecutionError(f"交叉校验缺少参数: {left_name}/{right_name}")
    unit = str(config.get("unit") or params[left_name].unit)
    left = _parameter_number(params[left_name], unit)
    right = _parameter_number(params[right_name], unit)
    tolerance = Decimal(str(config.get("tolerance") or "0"))
    comparator = str(config.get("operator") or "==")
    if comparator == "==" and tolerance > 0:
        passed = abs(left - right) <= tolerance
    else:
        passed = compare_values(left, comparator, right)
    message = f"{left_name}={left}{unit} {comparator} {right_name}={right}{unit}"
    return (
        passed,
        message,
        {
            left_name: _parameter_evidence(params[left_name], value=left, unit=unit),
            right_name: _parameter_evidence(params[right_name], value=right, unit=unit),
        },
        {"operator": comparator, "tolerance": str(tolerance), "unit": unit},
    )


def _latest_enabled_rules(db: Session, scheme_type_id: int) -> list[RuleDefinition]:
    rows = (
        db.query(RuleDefinition)
        .filter(RuleDefinition.scheme_type_id == scheme_type_id, RuleDefinition.enabled.is_(True))
        .order_by(RuleDefinition.rule_code, RuleDefinition.version.desc())
        .all()
    )
    latest: dict[str, RuleDefinition] = {}
    for row in rows:
        latest.setdefault(row.rule_code, row)
    return list(latest.values())


def _latest_formula(rule: RuleDefinition) -> FormulaDefinition | None:
    rows = [row for row in rule.formulas if row.enabled]
    rows.sort(key=lambda row: row.version, reverse=True)
    return rows[0] if rows else None


def execute_rules_for_task(db: Session, task: SchemeReviewTask) -> list[CalculationResult]:
    rules = _latest_enabled_rules(db, task.scheme_type_id)
    params = parameter_map(
        db.query(ParameterValue).filter(ParameterValue.task_id == task.id).all()
    )
    document_text: str | None = None
    results: list[CalculationResult] = []
    for rule in rules:
        config = _json_dict(rule.config_json)
        passed = False
        calculated: Decimal | None = None
        result_unit = ""
        inputs: dict[str, Any] = {}
        steps: list[dict[str, str]] = []
        expected: dict[str, Any] = {}
        message = ""
        error = ""
        formula_id: int | None = None
        try:
            if not rule.source_standard_no or not rule.source_clause or not rule.source_text:
                raise RuleExecutionError("规则缺少可追溯的规范编号、条款号或条款原文")
            if rule.rule_type == "required_section":
                if document_text is None:
                    document_text = extract_docx_text(task)
                passed, message, inputs = _evaluate_required_section(rule, config, document_text)
            elif rule.rule_type == "parameter_range":
                passed, calculated, result_unit, inputs, expected = _evaluate_parameter_range(
                    rule, config, params
                )
                message = str(expected.pop("message"))
            elif rule.rule_type == "cross_field":
                passed, message, inputs, expected = _evaluate_cross_field(config, params)
            elif rule.rule_type == "formula":
                formula = _latest_formula(rule)
                if formula is None:
                    raise RuleExecutionError("公式规则未配置启用公式")
                formula_id = formula.id
                variables_spec = _json_dict(formula.variables_json)
                variables: dict[str, Decimal] = {}
                for variable_name, raw_spec in variables_spec.items():
                    spec = raw_spec if isinstance(raw_spec, dict) else {"parameter": raw_spec}
                    parameter_name = str(spec.get("parameter") or variable_name)
                    if parameter_name not in params:
                        raise RuleExecutionError(f"缺少公式参数: {parameter_name}")
                    target_unit = str(spec.get("unit") or params[parameter_name].unit)
                    value = _parameter_number(params[parameter_name], target_unit)
                    variables[str(variable_name)] = value
                    inputs[str(variable_name)] = {
                        **_parameter_evidence(
                            params[parameter_name],
                            value=value,
                            unit=target_unit,
                        ),
                        "parameter": parameter_name,
                    }
                calculated, steps = safe_evaluate(formula.expression, variables)
                result_unit = formula.result_unit
                if formula.threshold_parameter:
                    threshold_row = params.get(formula.threshold_parameter)
                    if threshold_row is None:
                        raise RuleExecutionError(f"缺少阈值参数: {formula.threshold_parameter}")
                    threshold = _parameter_number(threshold_row, formula.result_unit)
                elif formula.threshold_value is not None:
                    threshold = Decimal(formula.threshold_value)
                else:
                    raise RuleExecutionError("公式未配置阈值")
                passed = compare_values(calculated, formula.comparator, threshold)
                expected = {"comparator": formula.comparator, "threshold": str(threshold), "unit": formula.result_unit, "expression": formula.expression}
                message = f"计算结果 {calculated}{formula.result_unit} {formula.comparator} {threshold}{formula.result_unit}"
            else:
                raise RuleExecutionError(f"未知规则类型: {rule.rule_type}")
        except Exception as exc:
            error = str(exc)
            message = f"规则无法完成确定性判定：{error}"
            passed = False
        row = CalculationResult(
            task_id=task.id,
            rule_id=rule.id,
            formula_id=formula_id,
            passed=passed,
            calculated_value=calculated,
            result_unit=result_unit,
            inputs_json=json.dumps(inputs, ensure_ascii=False, separators=(",", ":")),
            steps_json=json.dumps(steps, ensure_ascii=False, separators=(",", ":")),
            expected_json=json.dumps(expected, ensure_ascii=False, separators=(",", ":")),
            message=message,
            error_message=error,
            created_at=datetime.now(UTC),
        )
        db.add(row)
        results.append(row)
    db.flush()
    return results
