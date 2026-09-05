from io import BytesIO

from docx import Document

from app.services.paddle_document_parser import markdown_to_tree


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("工程概况", level=1)
    document.add_paragraph("脚手架搭设高度为 24m。")
    document.add_heading("施工工艺", level=1)
    document.add_paragraph("搭设流程说明。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_markdown_to_tree_preserves_paddle_order_and_word_anchors() -> None:
    tree = markdown_to_tree(
        "# 工程概况\n脚手架搭设高度为 24m。\n# 施工工艺\n搭设流程说明。",
        docx_bytes=_docx_bytes(),
    )
    nodes = tree["nodes"]
    assert [node["title"] for node in nodes] == ["工程概况", "施工工艺"]
    assert nodes[0]["content"] == ["脚手架搭设高度为 24m。"]
    assert nodes[0]["heading_para_index"] == 0
    assert nodes[1]["heading_para_index"] == 2


def test_markdown_to_tree_keeps_tables_and_formulas_as_content() -> None:
    tree = markdown_to_tree("# 计算书\n| 参数 | 数值 |\n| --- | --- |\n| H | 24m |\n$$N=Af$$")
    content = tree["nodes"][0]["content"]
    assert "| 参数 | 数值 |" in content
    assert "$$N=Af$$" in content


def test_markdown_to_tree_recognises_common_chinese_numbering() -> None:
    tree = markdown_to_tree(
        "第一章 工程概况\n1.1 项目概况\n正文内容\n一、施工部署\n（一）材料准备"
    )
    nodes = tree["nodes"]
    assert nodes[0]["title"] == "第一章 工程概况"
    assert nodes[0]["children"][0]["title"] == "1.1 项目概况"
    assert nodes[1]["title"] == "一、施工部署"
    assert nodes[1]["children"][0]["title"] == "（一）材料准备"
