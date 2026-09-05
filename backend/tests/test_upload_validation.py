from io import BytesIO
import zipfile

import pytest
from docx import Document

from app.services.upload_validation import UploadValidationError, validate_docx_upload


def _docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("危大工程专项施工方案", level=1)
    document.add_paragraph("测试正文")
    document.save(stream)
    return stream.getvalue()


def test_valid_docx_is_accepted() -> None:
    validate_docx_upload("方案.docx", _docx_bytes())


@pytest.mark.parametrize("filename", ["方案.pdf", "方案.doc", "方案"])
def test_non_docx_extension_is_rejected(filename: str) -> None:
    with pytest.raises(UploadValidationError, match="docx"):
        validate_docx_upload(filename, _docx_bytes())


def test_fake_or_incomplete_ooxml_is_rejected() -> None:
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
    with pytest.raises(UploadValidationError, match="OOXML"):
        validate_docx_upload("fake.docx", stream.getvalue())


def test_zip_path_traversal_is_rejected() -> None:
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", "<document />")
        archive.writestr("../outside.txt", "unsafe")
    with pytest.raises(UploadValidationError, match="不安全"):
        validate_docx_upload("unsafe.docx", stream.getvalue())
