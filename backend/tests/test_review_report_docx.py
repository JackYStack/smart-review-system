"""Tests for audit report Word generation."""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime

from docx import Document

from app.schemas.review_report import ReportIssue, ReportStep, ReviewReportV1
from app.services.review_report_docx import build_audit_report_docx


class _FakeScheme:
    category = "模板工程及支撑体系"
    name = "高大模板支撑"


class _FakeTask:
    original_filename = "睦邻中心高大支模安全专项方案测试.docx"
    status = "succeeded"
    finished_at = datetime.now(UTC)
    updated_at = datetime.now(UTC)
    review_conclusion = "issues_found"
    completeness_status = "partial"
    human_status = "pending"
    scheme_type = _FakeScheme()


_CHAPTER = "一、工程概况 > 1.模板支撑体系工程概况和特点"


def test_build_audit_report_docx_contains_sections_and_tables() -> None:
    report = ReviewReportV1(
        steps=[
            ReportStep(
                step_id="structure",
                passed=False,
                summary="共 1 项结构问题",
                issues=[
                    ReportIssue(
                        message="缺少应急预案章节",
                        severity="error",
                        anchor={"title_path": ["八、应急处置措施"]},
                        related={"kind": "missing_section"},
                    )
                ],
            ),
            ReportStep(
                step_id="compilation_basis",
                passed=False,
                summary="发现 1 条编制依据问题",
                issues=[
                    ReportIssue(
                        message="未引用现行规范",
                        severity="error",
                        evidence="无",
                        related={
                            "category": "现行缺失",
                            "standard_no": "GB51210-2016",
                            "doc_name": "建筑施工脚手架安全技术统一标准",
                            "suggestions": ["在编制依据中补充引用"],
                        },
                    )
                ],
            ),
            ReportStep(
                step_id="content",
                passed=False,
                summary="发现 2 条内容问题",
                issues=[
                    ReportIssue(
                        message="未明确高大支模范围内梁的跨度",
                        severity="error",
                        evidence="第3节高大支模部位表格中仅列出梁截面尺寸",
                        anchor={"title_path": _CHAPTER.split(" > ")},
                        related={
                            "suggestions": [
                                "在表格或正文中补充梁的跨度数据",
                            ],
                        },
                    ),
                    ReportIssue(
                        message="未描述周边环境情况",
                        severity="error",
                        evidence="第4节仅说明高大模板立杆支承在基础和楼板上",
                        anchor={"title_path": _CHAPTER.split(" > ")},
                        related={
                            "suggestions": [
                                "补充支撑地基的承载力、压实度等参数",
                            ],
                        },
                    ),
                ],
            ),
        ]
    )
    data = build_audit_report_docx(_FakeTask(), report, system_name="智能方案审核")
    assert len(data) > 1000

    doc = Document(io.BytesIO(data))
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "审核报告" in texts
    assert "结构审核" in texts
    assert "编制依据审核" in texts
    assert "内容审核" in texts
    assert "AI 辅助审查" in texts
    assert "本次并非完整审查" in texts
    assert "严重程度" not in texts

    heading3_styles = [
        p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 3"
    ]
    assert _CHAPTER in heading3_styles

    assert len(doc.tables) >= 4

    meta_table = doc.tables[0]
    assert meta_table.rows[0].cells[0].text == "项目"
    assert meta_table.rows[0].cells[1].text == "内容"
    meta_labels = [row.cells[0].text for row in meta_table.rows[1:]]
    assert "方案类型" in meta_labels
    assert "技术处理状态" in meta_labels
    assert "AI审查结论" in meta_labels
    assert "审查完整性" in meta_labels
    assert "人工复核状态" in meta_labels
    assert "摘要" in meta_labels
    summary_row = next(row for row in meta_table.rows[1:] if row.cells[0].text == "摘要")
    assert "本次审核任务状态为" in summary_row.cells[1].text
    assert "摘要：" not in summary_row.cells[1].text

    content_table = next(
        t for t in doc.tables if t.rows[0].cells[0].text == "序号"
    )
    assert content_table.rows[0].cells[1].text == "证据状态"
    assert content_table.rows[0].cells[2].text == "问题"
    assert content_table.rows[0].cells[3].text == "方案原文/验算过程"
    assert content_table.rows[0].cells[4].text == "法规依据"
    assert content_table.rows[0].cells[5].text == "修改建议"
    assert "待人工" in content_table.rows[1].cells[1].text
    assert len(content_table.rows) == 3


def test_parse_review_report_json_roundtrip() -> None:
    from app.api.review_tasks import _audit_report_filename, _parse_review_report_json

    assert _audit_report_filename("test方案.docx") == "test方案_审核报告.docx"
    raw = json.dumps({"version": 1, "steps": [{"step_id": "structure", "passed": True, "issues": []}]})
    parsed = _parse_review_report_json(raw)
    assert parsed is not None
    assert parsed.steps[0].step_id == "structure"
